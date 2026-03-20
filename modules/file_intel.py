"""
File Intelligence Module — ported/inspired by ClatScope Mini (Clats97)
Extracts metadata from images, PDFs, Office docs, audio files.
Also computes MD5 / SHA1 / SHA256 / SHA512 hashes.
"""
from __future__ import annotations

import hashlib
import os
import socket
import stat
from datetime import datetime
from pathlib import Path

from .utils import console, print_section, print_result


def _fmt_size(b: int) -> str:
    for unit in ("B", "KB", "MB", "GB"):
        if b < 1024:
            return f"{b:.1f} {unit}"
        b /= 1024
    return f"{b:.1f} TB"


def _fmt_ts(ts: float) -> str:
    return datetime.fromtimestamp(ts).strftime("%Y-%m-%d %H:%M:%S")


# ── File Hash ─────────────────────────────────────────────────────────────────

def compute_hashes(file_path: str) -> dict:
    try:
        with open(file_path, "rb") as f:
            data = f.read()
        return {
            "md5":    hashlib.md5(data).hexdigest(),
            "sha1":   hashlib.sha1(data).hexdigest(),
            "sha256": hashlib.sha256(data).hexdigest(),
            "sha512": hashlib.sha512(data).hexdigest(),
        }
    except Exception as exc:
        return {"error": str(exc)}


# ── Base file info ────────────────────────────────────────────────────────────

def _base_info(path: str) -> dict:
    st = os.stat(path)
    return {
        "path":     path,
        "size":     _fmt_size(st.st_size),
        "created":  _fmt_ts(st.st_ctime),
        "modified": _fmt_ts(st.st_mtime),
        "accessed": _fmt_ts(st.st_atime),
    }


# ── Image (EXIF + GPS) ────────────────────────────────────────────────────────

def _read_image(path: str) -> dict:
    from PIL import Image
    from PIL.ExifTags import TAGS, GPSTAGS

    info = _base_info(path)
    try:
        img = Image.open(path)
        info["format"] = img.format or ""
        info["mode"]   = img.mode  or ""
        info["size_px"] = f"{img.width}x{img.height}"

        exif_raw = img._getexif() if hasattr(img, "_getexif") else None
        if exif_raw:
            exif: dict = {}
            gps:  dict = {}
            for tag_id, value in exif_raw.items():
                tag = TAGS.get(tag_id, str(tag_id))
                if tag == "GPSInfo" and isinstance(value, dict):
                    for gtag_id, gvalue in value.items():
                        gps[GPSTAGS.get(gtag_id, gtag_id)] = str(gvalue)
                else:
                    exif[tag] = str(value)[:120]
            info["exif"] = exif
            if gps:
                info["gps"] = gps
    except Exception as exc:
        info["exif_error"] = str(exc)
    return info


# ── PDF ───────────────────────────────────────────────────────────────────────

def _read_pdf(path: str) -> dict:
    import PyPDF2
    info = _base_info(path)
    try:
        with open(path, "rb") as f:
            reader = PyPDF2.PdfReader(f)
        info["pages"] = len(reader.pages)
        meta = reader.metadata or {}
        for k, v in meta.items():
            clean_key = k.lstrip("/")
            info[clean_key] = str(v)[:200]
    except Exception as exc:
        info["pdf_error"] = str(exc)
    return info


# ── Office (docx / xlsx / pptx) ──────────────────────────────────────────────

def _read_docx(path: str) -> dict:
    import docx as _docx
    info = _base_info(path)
    try:
        doc  = _docx.Document(path)
        core = doc.core_properties
        for attr in ("author", "created", "modified", "last_modified_by",
                     "revision", "title", "subject", "keywords", "description"):
            val = getattr(core, attr, None)
            if val:
                info[attr] = str(val)[:200]
        info["paragraphs"] = len(doc.paragraphs)
    except Exception as exc:
        info["docx_error"] = str(exc)
    return info


def _read_xlsx(path: str) -> dict:
    import openpyxl
    info = _base_info(path)
    try:
        wb   = openpyxl.load_workbook(path, read_only=True, data_only=True)
        props = wb.properties
        for attr in ("creator", "lastModifiedBy", "created", "modified",
                     "title", "subject", "keywords", "description"):
            val = getattr(props, attr, None)
            if val:
                info[attr] = str(val)[:200]
        info["sheets"] = wb.sheetnames
    except Exception as exc:
        info["xlsx_error"] = str(exc)
    return info


def _read_pptx(path: str) -> dict:
    from pptx import Presentation
    info = _base_info(path)
    try:
        prs  = Presentation(path)
        core = prs.core_properties
        for attr in ("author", "created", "modified", "last_modified_by",
                     "revision", "title", "subject", "keywords"):
            val = getattr(core, attr, None)
            if val:
                info[attr] = str(val)[:200]
        info["slides"] = len(prs.slides)
    except Exception as exc:
        info["pptx_error"] = str(exc)
    return info


# ── Audio ─────────────────────────────────────────────────────────────────────

def _read_audio(path: str) -> dict:
    from tinytag import TinyTag
    info = _base_info(path)
    try:
        tag = TinyTag.get(path)
        for attr in ("title", "artist", "album", "year", "genre",
                     "duration", "bitrate", "samplerate", "channels"):
            val = getattr(tag, attr, None)
            if val is not None:
                info[attr] = str(round(val, 2)) if isinstance(val, float) else str(val)
    except Exception as exc:
        info["audio_error"] = str(exc)
    return info


# ── Dispatcher ────────────────────────────────────────────────────────────────

_EXT_MAP = {
    ".jpg": _read_image, ".jpeg": _read_image, ".png":  _read_image,
    ".gif": _read_image, ".bmp":  _read_image, ".tiff": _read_image,
    ".webp":_read_image,
    ".pdf": _read_pdf,
    ".docx":_read_docx, ".doc": _read_docx,
    ".xlsx":_read_xlsx, ".xls": _read_xlsx,
    ".pptx":_read_pptx, ".ppt": _read_pptx,
    ".mp3": _read_audio, ".flac": _read_audio, ".ogg": _read_audio,
    ".m4a": _read_audio, ".wav":  _read_audio, ".aac": _read_audio,
}


def read_file_metadata(file_path: str) -> dict:
    ext = Path(file_path).suffix.lower()
    reader = _EXT_MAP.get(ext)
    if reader:
        return reader(file_path)
    # Generic fallback
    return _base_info(file_path)


# ── Public run() ──────────────────────────────────────────────────────────────

def run(file_path: str, show_hashes: bool = True) -> None:
    print_section("FILE INTELLIGENCE")

    if not os.path.isfile(file_path):
        console.print(f"  [red]File not found:[/red] {file_path}")
        return

    console.print(f"  [dim]File:[/dim] [bold]{file_path}[/bold]\n")

    meta = read_file_metadata(file_path)
    for key, val in meta.items():
        if key == "exif":
            console.print(f"  [bold cyan]EXIF data:[/bold cyan]")
            for ek, ev in val.items():
                if ek not in ("MakerNote", "UserComment", "PrintImageMatching"):
                    console.print(f"    [dim]{ek}:[/dim] {ev}")
        elif key == "gps":
            console.print(f"  [bold cyan]GPS data:[/bold cyan]")
            for gk, gv in val.items():
                console.print(f"    [dim]{gk}:[/dim] {gv}")
        elif key == "sheets":
            print_result("Sheets", ", ".join(val))
        elif not key.endswith("_error") and key != "path":
            print_result(key.replace("_", " ").title(), str(val))

    if show_hashes:
        console.print()
        console.print("  [bold cyan]Hashes:[/bold cyan]")
        hashes = compute_hashes(file_path)
        for algo, val in hashes.items():
            if algo != "error":
                console.print(f"    [dim]{algo.upper()}:[/dim] {val}")
